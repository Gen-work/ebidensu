"""Pre-delivery self-review of J4 evidence workbooks (GFIX receive).

Checks the patterns 加瀬さん's review (Check Sheet_J4) has sent back, and the
ones the 10/23 self-review found (docs/ebi-dance/KNOWLEDGE.md section 5).
Read-only: never writes the workbook.

Usage:
  python3 check_book.py <book.xlsx> [<book.xlsx> ...]

Prints one block per book: NG lines (fix before delivery), WARN lines (look),
then a one-line summary. Exit code 1 when any book has an NG.

What it does NOT check (look at the pictures, drawings.py render):
  - whether a red box sits on the right row of a screenshot
  - whether the GIFT list box / Jenkins box / DF box match each other
  - whether a DF that differs carries an explanation (callout)
"""
import re
import sys
import warnings
import zipfile
import xml.etree.ElementTree as ET

import openpyxl

sys.path.insert(0, __import__('os').path.dirname(__file__))
from drawings import Book, sheet_items, M  # noqa: E402

warnings.filterwarnings('ignore')

S_SEND = '送信データ'
S_GIFT_RECV = 'GIFT受信結果'
S_GFIX_RECV = 'GFIX受信結果'
S_COMPARE = 'GIFTデータvsGFIXデータ'
RECV_SHEETS = (S_GIFT_RECV, S_GFIX_RECV)
FONT = 'ＭＳ ゴシック'
YELLOW = 'FFFF00'
# 'GIFTReceiverUnzip START <file>' / 'GFIXReceiver END <file>'
RUN_LINE = re.compile(r'^(GIFT|GFIX)Receiver(Unzip)? (START|END) (\S+)')
MOJIBAKE = re.compile('[�縺繧繝譁蜿驕荳隱]|[\x80-\x9f]')
CHARS_PER_COL = 3.0     # log rows: ~3 characters per column at the sheet's widths


class Report:
    def __init__(self, name):
        self.name = name
        self.lines = []

    def ng(self, where, msg, ref=''):
        self.lines.append(('NG', where, msg, ref))

    def warn(self, where, msg, ref=''):
        self.lines.append(('WARN', where, msg, ref))

    def print(self):
        print('=' * 72)
        print(self.name)
        for lvl, where, msg, ref in self.lines:
            print('  %-4s %-28s %s%s' % (lvl, where, msg, ('  (%s)' % ref) if ref else ''))
        n_ng = sum(1 for l in self.lines if l[0] == 'NG')
        n_w = len(self.lines) - n_ng
        print('  -> %s: %d NG, %d WARN' % ('NG' if n_ng else 'OK', n_ng, n_w))
        return n_ng


def is_yellow(c):
    f = c.fill
    return (f is not None and f.fill_type == 'solid'
            and isinstance(f.fgColor.rgb, str) and f.fgColor.rgb.upper().endswith(YELLOW))


def check_view(path, rep):
    """First sheet active; every sheet scrolled to the top with the cursor on A3
    (Check Sheet_J4: 'セル選択をA3にした上で保存' / '保存時のカーソル位置とシートが違います')."""
    bk = Book(path)
    for i, (name, part, state) in enumerate(bk.sheets()):
        if state != 'visible':
            continue
        sv = bk.xml(part).find(M + 'sheetViews')
        v = sv.find(M + 'sheetView') if sv is not None else None
        if v is None:
            continue
        if (v.get('tabSelected') == '1') != (i == 0):
            rep.ng(name, 'active sheet on save must be the first sheet' if i == 0
                   else 'this sheet is selected; save with the first sheet active')
        pane = v.find(M + 'pane')
        top = pane.get('topLeftCell') if pane is not None else v.get('topLeftCell', 'A1')
        if top not in ('A1', 'A3'):
            rep.ng(name, 'scrolled to %s; scroll to the top before saving' % top)
        sels = v.findall(M + 'selection')
        act = [s.get('activeCell') for s in sels if s.get('pane') in (None, 'bottomLeft')]
        if act and act[-1] not in ('A1', 'A3'):
            rep.ng(name, 'cursor on %s; put it on A3 before saving' % act[-1])


def check_fonts(wb, rep):
    for sn in (S_SEND, S_GIFT_RECV, S_GFIX_RECV, S_COMPARE):
        if sn not in wb.sheetnames:
            rep.ng(sn, 'sheet missing')
            continue
        bad = {}
        for row in wb[sn].iter_rows():
            for c in row:
                if c.value is not None and c.font.name != FONT:
                    bad.setdefault(c.font.name, c.coordinate)
        for font, where in bad.items():
            rep.ng('%s!%s' % (sn, where), 'font %r (should be %s)' % (font, FONT), 'Check Sheet: 標準化違反')


def check_recv_logs(wb, rep):
    """Log blocks on GIFT/GFIX受信結果: the label above each block names the log it
    was cut from; each block has whole runs (START ... END + date line) and nothing
    of another run."""
    for sn in RECV_SHEETS:
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        rows = [(r, str(ws.cell(r, 2).value)) for r in range(1, ws.max_row + 1)
                if ws.cell(r, 2).value is not None]
        labels = [i for i, (r, v) in enumerate(rows) if v.strip().endswith('.log')]
        for k, i in enumerate(labels):
            r_label, label = rows[i]
            end_i = labels[k + 1] if k + 1 < len(labels) else len(rows)
            block = rows[i + 1:end_i]
            where = '%s!B%d' % (sn, r_label)
            if not ws.cell(r_label, 2).font.b:
                rep.ng(where, 'log label %r is not bold' % label, 'Check Sheet 10/12')
            runs = [(r, RUN_LINE.match(v.strip())) for r, v in block]
            runs = [(r, m) for r, m in runs if m]
            if not runs:
                continue
            side, unzip = runs[0][1].group(1), runs[0][1].group(2)
            expect = '%sReceive%s.log' % (side, 'Unzip' if unzip else '')
            if label.strip() != expect:
                rep.ng(where, 'label %r but the lines are %sReceiver%s -> %r'
                       % (label, side, unzip or '', expect), 'K-36')
            if block and not RUN_LINE.match(block[0][1].strip()):
                rep.ng('%s!B%d' % (sn, block[0][0]), 'block does not start with a START line: %r'
                       % block[0][1][:50], 'K-37')
            starts = {}
            for r, m in runs:
                if m.group(3) != 'START':
                    continue
                if m.group(4) in starts:
                    rep.ng('%s!B%d' % (sn, r), 'START %s again (B%d): two START lines with one file name'
                           % (m.group(4), starts[m.group(4)]), 'K-37')
                else:
                    starts[m.group(4)] = r
            ends = {m.group(4): r for r, m in runs if m.group(3) == 'END'}
            for f, r in ends.items():
                if f not in starts:
                    rep.ng('%s!B%d' % (sn, r), 'END %s has no START line in this block' % f, 'K-37')
            for f, r in starts.items():
                if f not in ends:
                    rep.ng('%s!B%d' % (sn, r), 'START %s has no END line: a line of another run?' % f, 'K-37')
            last_end = max(ends.values()) if ends else None
            if last_end:
                tail = [(r, v) for r, v in block if r > last_end]
                if len(tail) > 1:
                    rep.ng('%s!B%d' % (sn, tail[1][0]), '%d line(s) after the last END + date line: %r'
                           % (len(tail) - 1, tail[1][1][:50]), 'K-37')
        for r, v in rows:
            if MOJIBAKE.search(v):
                rep.ng('%s!B%d' % (sn, r), 'mojibake? %r' % v[:60], 'Check Sheet 10/21')


def check_yellow(wb, rep):
    """Yellow belongs on a log line and runs to the end of its text."""
    for sn in RECV_SHEETS:
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        for row in ws.iter_rows():
            ys = [c for c in row if is_yellow(c)]
            if not ys:
                continue
            texts = [c for c in row if c.value is not None]
            r = row[0].row
            if not texts:
                rep.ng('%s!%s' % (sn, ys[0].coordinate), 'yellow fill on an empty row (often under a picture)', 'K-38')
                continue
            n = len(str(texts[0].value).rstrip())
            cols = max(c.column for c in ys) - min(c.column for c in ys) + 1
            if cols * CHARS_PER_COL < n - 2 * CHARS_PER_COL:
                rep.warn('%s!B%d' % (sn, r), 'yellow stops before the text ends (%d cols for %d chars)' % (cols, n),
                         'Check Sheet 10/20')
            elif cols * CHARS_PER_COL > n + 4 * CHARS_PER_COL:
                rep.warn('%s!B%d' % (sn, r), 'yellow runs past the text (%d cols for %d chars)' % (cols, n),
                         'Check Sheet 07/01')


def compare_label_rows(path, wb):
    """GIFTデータvsGFIXデータ label rows (a correl id in B): [(row, vals, stale)].
    A set's label row is row 3 or the row right above a picture anchored in
    column B; a label anywhere else was left by an earlier run whose pictures
    sat elsewhere (10/27 F/K/R/S: B65, B66, B71 -- K-48)."""
    cmp_ = wb[S_COMPARE]
    bk = Book(path)
    part = next((p for n, p, st in bk.sheets() if n == S_COMPARE), None)
    pic_rows = {it['anchor'][1] + 1 for it in sheet_items(bk, part)
                if it['kind'] == 'pic' and it.get('anchor') and it['anchor'][0] == 1} if part else set()
    out = []
    for r in range(3, cmp_.max_row + 1):
        vals = {c.column_letter: str(c.value).strip() for c in cmp_[r] if c.value is not None}
        if vals.get('B') and re.match(r'^[A-Z0-9]{6,}$', vals['B']):
            out.append((r, vals, bool(pic_rows) and r != 3 and (r + 1) not in pic_rows))
    return out


def check_compare_labels(path, wb, rep):
    """GIFTデータvsGFIXデータ: each file pair is labelled correl id in B and the STEP
    name of 送信データ's matching label row in G (Check Sheet 10/13, K-32); no
    label left where no set starts (K-48)."""
    if S_SEND not in wb.sheetnames or S_COMPARE not in wb.sheetnames:
        return
    send = wb[S_SEND]
    send_labels = []
    for r in range(3, send.max_row + 1):
        a = send.cell(r, 1).value
        if a and re.match(r'^[A-Z0-9]{6,}$', str(a).strip()):
            send_labels.append((str(a).strip(), (str(send.cell(r, 2).value or '')).strip()))
    cmp_rows = []
    for r, vals, stale in compare_label_rows(path, wb):
        if stale:
            cells = '/'.join('%s%d' % (c, r) for c in ('B', 'G') if vals.get(c))
            rep.ng('%s!B%d' % (S_COMPARE, r), '%s: a label on no set\'s row (no picture starts at row %d) '
                   '-- left by an earlier run, delete it' % (cells, r + 1), 'K-48')
            continue
        cmp_rows.append((r, vals))
    if len(cmp_rows) != len(send_labels):
        rep.ng(S_COMPARE, '%d label row(s) but 送信データ has %d file group(s)'
               % (len(cmp_rows), len(send_labels)), 'K-39')
    for k, (r, vals) in enumerate(cmp_rows):
        if k >= len(send_labels):
            break
        correl, step = send_labels[k]
        if vals['B'] != correl:
            rep.ng('%s!B%d' % (S_COMPARE, r), 'correl %r, 送信データ group %d says %r' % (vals['B'], k + 1, correl), 'K-39')
        if not step:
            continue
        if vals.get('G') != step:
            elsewhere = [col for col, v in vals.items() if v == step and col != 'G']
            msg = ('STEP name in %s%d, should be in G%d' % (elsewhere[0], r, r)) if elsewhere \
                else 'G%d is %r, 送信データ says %r' % (r, vals.get('G', ''), step)
            rep.ng('%s!G%d' % (S_COMPARE, r), msg, 'K-39')


UPLOAD = re.compile(r"Uploading '[^']*' \(([0-9,]+) bytes\) to ")


def gfix_job_blocks(ws):
    """GFIX受信ログ on GFIX受信結果: one block per job log (split on 'Start Date
    and Time'), each @{ row; bytes; renames = [(row, yellow)] }."""
    rows = [(r, str(ws.cell(r, 2).value)) for r in range(1, ws.max_row + 1)
            if ws.cell(r, 2).value is not None]
    start = next((i for i, (r, v) in enumerate(rows) if v.strip() == 'GFIX受信ログ'), None)
    if start is None:
        return []
    blocks = []
    for r, v in rows[start + 1:]:
        if v.strip().endswith('.log'):
            break
        if 'Start Date and Time' in v:
            blocks.append({'row': r, 'bytes': None, 'renames': []})
        if not blocks:
            continue
        m = UPLOAD.search(v)
        if m and blocks[-1]['bytes'] is None:
            blocks[-1]['bytes'] = int(m.group(1).replace(',', ''))
        if re.search(r"INFO\s+Renaming '", v):
            blocks[-1]['renames'].append((r, any(is_yellow(ws.cell(r, c)) for c in range(2, 6))))
    return blocks


def compare_groups(path, wb):
    """GIFTデータvsGFIXデータ: per file set (a correl-id label row in B opens one),
    whether its DF is an empty file -- only the 同一内容 box, no line-number box."""
    if S_COMPARE not in wb.sheetnames:
        return []
    starts = [r for r, vals, stale in compare_label_rows(path, wb) if not stale]
    if not starts:
        return []
    bk = Book(path)
    part = next((p for n, p, st in bk.sheets() if n == S_COMPARE), None)
    boxes = [it['anchor'][1] + 1 for it in sheet_items(bk, part)
             if it['kind'] == 'sp' and (it.get('line') or '').upper() == 'FF0000' and it.get('anchor')]
    out = []
    for k, r0 in enumerate(starts):
        r1 = starts[k + 1] if k + 1 < len(starts) else 10 ** 9
        n = sum(1 for b in boxes if r0 <= b < r1)
        out.append({'row': r0, 'boxes': n, 'empty': n == 1})
    return out


def check_gfix_job_logs(path, wb, rep):
    """GFIX受信ログ (加瀬さん 10/26, C/F/MJODWAF1): the Renaming line of each job
    log is yellow (K-45); with several files the job logs follow the
    送信データ / DF order -- the empty file (smallest upload, 99 bytes) sits
    where the DF set is empty (K-44)."""
    if S_GFIX_RECV not in wb.sheetnames:
        return
    blocks = gfix_job_blocks(wb[S_GFIX_RECV])
    for b in blocks:
        for r, yellow in b['renames']:
            if not yellow:
                rep.ng('%s!B%d' % (S_GFIX_RECV, r), 'GFIX受信ログ: the Renaming line is not yellow', 'K-45')
    if len(blocks) < 2:
        return
    groups = compare_groups(path, wb)
    sizes = [b['bytes'] for b in blocks]
    if len(groups) != len(blocks) or None in sizes or len(set(sizes)) != len(sizes):
        return
    empty = {k for k, g in enumerate(groups) if g['empty']}
    if not empty or len(empty) == len(groups):
        return
    smallest = set(sorted(range(len(sizes)), key=lambda k: sizes[k])[:len(empty)])
    if smallest != empty:
        order = ', '.join('B%d %s bytes' % (b['row'], format(b['bytes'], ',')) for b in blocks)
        want = ', '.join('set %d %s' % (k + 1, 'empty' if g['empty'] else 'data') for k, g in enumerate(groups))
        rep.ng('%s!B%d' % (S_GFIX_RECV, blocks[0]['row']),
               'GFIX受信ログ order differs from the DF sets (%s; DF: %s)' % (order, want), 'K-44')


def overlap(a, b):
    return max(0, min(a[2], b[2]) - max(a[0], b[0])) * max(0, min(a[3], b[3]) - max(a[1], b[1]))


def check_shapes(path, rep):
    """Every red box should sit on a picture; a box over nothing is a leftover.
    A picture pasted by hand over one the tool placed (verifyPic_*) leaves the
    old picture and its boxes (verifyMark_*) underneath, the boxes measured on
    the old one (10/27 LJODWAF1: 図 1 over verifyPic_df_4, the 同一内容 box
    no longer on 同一内容 -- K-49)."""
    bk = Book(path)
    for name, part, state in bk.sheets():
        items = sheet_items(bk, part)
        for t in items:
            if t['kind'] != 'pic' or not t['name'].startswith('verifyPic_'):
                continue
            area = max(1, (t['box'][2] - t['box'][0]) * (t['box'][3] - t['box'][1]))
            for h in items:
                if h['kind'] == 'pic' and not h['name'].startswith('verify') and h.get('z', 0) > t.get('z', 0) \
                        and overlap(h['box'], t['box']) >= 0.5 * area:
                    a = t.get('anchor')
                    rep.ng('%s!%s' % (name, ('B%d' % (a[1] + 1)) if a else '?'),
                           '%r is pasted over %r: the old picture and its boxes (verifyMark_%s_*) stay underneath, '
                           'measured on the old one -- delete the old picture, re-place the boxes on the new one'
                           % (h['name'], t['name'], t['name'][len('verifyPic_'):]), 'K-49')
        pics = [i['box'] for i in items if i['kind'] == 'pic']
        for it in items:
            if it['kind'] != 'sp' or (it.get('line') or '').upper() != 'FF0000':
                continue
            b = it['box']
            if not any(b[0] < p[2] and b[2] > p[0] and b[1] < p[3] and b[3] > p[1] for p in pics):
                a = it.get('anchor')
                rep.ng(name, 'red box %r is not on any picture (anchor row %s)'
                       % (it['name'], a[1] + 1 if a else '?'), 'Check Sheet 08/04')


def check_book(path):
    rep = Report(path)
    try:
        wb = openpyxl.load_workbook(path)
    except (zipfile.BadZipFile, OSError) as ex:
        rep.ng('-', 'cannot open: %s' % ex)
        return rep.print()
    check_view(path, rep)
    check_fonts(wb, rep)
    check_recv_logs(wb, rep)
    check_yellow(wb, rep)
    check_compare_labels(path, wb, rep)
    check_gfix_job_logs(path, wb, rep)
    check_shapes(path, rep)
    return rep.print()


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    bad = sum(1 for p in sys.argv[1:] if check_book(p))
    print('=' * 72)
    print('%d book(s), %d with NG' % (len(sys.argv) - 1, bad))
    sys.exit(1 if bad else 0)
